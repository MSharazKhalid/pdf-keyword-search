# search_pdf_pages.py
import os
import json
import fitz  # install PyMuPDF Not Fitz
import tkinter as tk
from tkinter import simpledialog, filedialog
from openpyxl import Workbook, load_workbook


def ask_for_file(title="Select PDF", filetypes=(("PDF files", "*.pdf"), ("All files", "*.*"))):
    root = tk.Tk()
    root.withdraw()
    path = filedialog.askopenfilename(
        title=title,
        filetypes=filetypes
    )
    root.destroy()
    return path.strip() if path else ""


def ask_for_query():
    root = tk.Tk()
    root.withdraw()
    query = simpledialog.askstring("Find in PDF", "Enter search text")
    root.destroy()
    return query.strip() if query else ""


def ask_for_mode():
    root = tk.Tk()
    root.withdraw()
    mode = simpledialog.askstring(
        "Choose mode",
        "Type 1 for single search\nType 2 for Excel column search"
    )
    root.destroy()
    return mode.strip() if mode else ""


def ask_for_sheet_name():
    root = tk.Tk()
    root.withdraw()
    sheet_name = simpledialog.askstring("Excel sheet", "Enter sheet name")
    root.destroy()
    return sheet_name.strip() if sheet_name else ""


def ask_for_column_ref():
    root = tk.Tk()
    root.withdraw()
    col_ref = simpledialog.askstring(
        "Excel column",
        "Enter column (like A, B, C, or 1, 2, 3)"
    )
    root.destroy()
    return col_ref.strip() if col_ref else ""


def column_ref_to_index(ref):
    ref = ref.strip()
    if not ref:
        return None

    if ref.isdigit():
        return int(ref)

    ref = ref.upper()
    col_index = 0
    for ch in ref:
        if not ("A" <= ch <= "Z"):
            return None
        col_index = col_index * 26 + (ord(ch) - ord("A") + 1)
    return col_index


def index_pdf(pdf_path):
    doc = fitz.open(pdf_path)
    pages_text = []
    for i in range(len(doc)):
        try:
            page = doc.load_page(i)
            text = page.get_text("text")
            pages_text.append(text.casefold())
        except Exception:
            pages_text.append("")
    doc.close()
    return pages_text


def search_term_in_index(pages_text, term):
    needle = str(term).casefold()
    if not needle:
        return []

    found_pages = []
    for idx, hay in enumerate(pages_text):
        if needle in hay:
            found_pages.append(idx + 1)  # 1 based page numbers
    return found_pages


def search_pdf(pdf_path, query):
    results = []
    doc = fitz.open(pdf_path)
    needle = query.casefold()

    for i in range(len(doc)):
        try:
            page = doc.load_page(i)
            text = page.get_text("text")
            hay = text.casefold()
            if needle in hay:
                count = hay.count(needle)
                results.append((i + 1, count))
        except Exception:
            continue

    doc.close()
    return results


def save_results(pdf_path, query, hits):
    out = {
        "pdf": os.path.abspath(pdf_path),
        "query": query,
        "total_pages_with_hits": len(hits),
        "pages": [{"page": p, "hits": c} for p, c in hits]
    }

    base = os.path.splitext(os.path.basename(pdf_path))[0]
    out_name = f"{base}_search_results.json"
    with open(out_name, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    return os.path.abspath(out_name)


def print_summary(hits):
    if not hits:
        print("No matches found.")
        return
    pages_only = [str(p) for p, _ in hits]
    print("Match found on pages:")
    print(", ".join(pages_only))
    print(f"Total pages with matches: {len(hits)}")
    total_hits = sum(c for _, c in hits)
    print(f"Total keyword occurrences: {total_hits}")


def load_excel_column_values(excel_path, sheet_name, col_index):
    wb = load_workbook(excel_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in Excel file.")
    ws = wb[sheet_name]

    values = []
    for row in range(2, ws.max_row + 1):  # skip header row
        cell_value = ws.cell(row=row, column=col_index).value
        if cell_value is None:
            continue
        text = str(cell_value).strip()
        if not text:
            continue
        values.append(text)
    return values


def save_excel_results(excel_path, values, results):
    base = os.path.splitext(os.path.basename(excel_path))[0]
    out_name = f"{base}_pdf_search_results.xlsx"
    out_full_path = os.path.join(os.path.dirname(os.path.abspath(excel_path)), out_name)

    wb = Workbook()
    ws = wb.active
    ws.title = "SearchResults"

    ws.cell(row=1, column=1, value="Value")
    ws.cell(row=1, column=2, value="Status")
    ws.cell(row=1, column=3, value="Pages")

    row_idx = 2
    for value in values:
        pages = results.get(value, [])
        status = "Found" if pages else "Not found"
        pages_str = ", ".join(str(p) for p in pages) if pages else ""

        ws.cell(row=row_idx, column=1, value=value)
        ws.cell(row=row_idx, column=2, value=status)
        ws.cell(row=row_idx, column=3, value=pages_str)
        row_idx += 1

    wb.save(out_full_path)
    return out_full_path


def excel_search_workflow(pdf_path):
    excel_path = ask_for_file(
        title="Select Excel file",
        filetypes=(("Excel files", "*.xlsx;*.xlsm;*.xltx;*.xltm"), ("All files", "*.*"))
    )
    if not excel_path:
        print("No Excel selected.")
        return

    sheet_name = ask_for_sheet_name()
    if not sheet_name:
        print("No sheet name entered.")
        return

    col_ref = ask_for_column_ref()
    if not col_ref:
        print("No column specified.")
        return

    col_index = column_ref_to_index(col_ref)
    if not col_index or col_index < 1:
        print("Invalid column.")
        return

    try:
        values = load_excel_column_values(excel_path, sheet_name, col_index)
    except Exception as e:
        print(f"Error reading Excel: {e}")
        return

    if not values:
        print("No values found in the selected column.")
        return

    print("Indexing PDF, please wait...")
    pages_text = index_pdf(pdf_path)
    print("Searching values in PDF...")

    results = {}
    for value in values:
        pages = search_term_in_index(pages_text, value)
        results[value] = pages

    out_excel = save_excel_results(excel_path, values, results)
    print(f"Excel results saved to: {out_excel}")


def main():
    pdf_path = ask_for_file()
    if not pdf_path:
        print("No PDF selected.")
        return

    mode = ask_for_mode()

    if mode == "2":
        excel_search_workflow(pdf_path)
        return

    query = ask_for_query()
    if not query:
        print("No search text entered.")
        return

    print("Scanning PDF, please wait...")
    hits = search_pdf(pdf_path, query)
    print_summary(hits)
    out_file = save_results(pdf_path, query, hits)
    print(f"Results saved to: {out_file}")


if __name__ == "__main__":
    main()
